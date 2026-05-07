"""Create notebook, README, demo script, and PowerPoint deck from pipeline outputs."""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import nbformat as nbf
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUTPUTS = ROOT / "outputs"
CHARTS = OUTPUTS / "charts"
DECK = ROOT / "deck" / "Transcript_Intelligence_Presentation.pptx"
NOTEBOOK = ROOT / "notebooks" / "01_transcript_intelligence_analysis.ipynb"
README = ROOT / "README.md"
DEMO = ROOT / "demo_script.md"

NAVY = RGBColor(23, 32, 51)
BLUE = RGBColor(45, 108, 223)
TEAL = RGBColor(32, 163, 158)
CORAL = RGBColor(242, 95, 92)
AMBER = RGBColor(245, 184, 65)
GRAY = RGBColor(122, 134, 154)
LIGHT = RGBColor(245, 247, 250)
WHITE = RGBColor(255, 255, 255)


def load_outputs():
    processed = pd.read_csv(OUTPUTS / "processed_transcripts.csv")
    topics = pd.read_csv(OUTPUTS / "topic_summary.csv")
    sentiment = pd.read_csv(OUTPUTS / "sentiment_summary.csv")
    escalation = pd.read_csv(OUTPUTS / "escalation_risks.csv")
    features = pd.read_csv(OUTPUTS / "feature_requests.csv")
    churn = pd.read_csv(OUTPUTS / "churn_risks.csv")
    stakeholder = pd.read_csv(OUTPUTS / "stakeholder_views.csv")
    metadata = json.loads((OUTPUTS / "run_metadata.json").read_text(encoding="utf-8"))
    return processed, topics, sentiment, escalation, features, churn, stakeholder, metadata


def write_markdown_assets():
    processed, topics, sentiment, escalation, features, churn, stakeholder, metadata = load_outputs()
    top_theme = topics.iloc[0]
    negative_support = int(sentiment.loc[sentiment["call_type"] == "support", "negative_transcripts"].iloc[0])
    support_total = int(sentiment.loc[sentiment["call_type"] == "support", "transcripts"].iloc[0])
    readme = f"""# Transcript Intelligence

This project is a take-home prototype for a B2B Enterprise SaaS product called **Transcript Intelligence**. It turns raw call transcripts into product, support, account, and engineering insights.

## Objective

Build an end-to-end pipeline that ingests transcript folders, normalizes the data, categorizes transcripts by business theme, scores sentiment, extracts additional insights, and packages the findings for a 30-minute leadership presentation.

## Key Findings

- Processed **{metadata['transcripts']} transcripts**: {metadata['call_type_counts'].get('support', 0)} support, {metadata['call_type_counts'].get('external', 0)} external/customer, and {metadata['call_type_counts'].get('internal', 0)} internal calls.
- Identified **{metadata['themes']} major themes**. The largest theme was **{top_theme['theme']}** with **{int(top_theme['transcript_count'])} transcripts**.
- Support calls carried the clearest risk signal: **{negative_support} of {support_total} support transcripts** were negative, concentrated around technical reliability, outage, SLA, and escalation language.
- External calls were more positive overall, but renewal/pricing conversations still surfaced trust, outage, and competitive risk signals.
- Additional insight layers identify escalation risk, feature requests, churn/renewal risk, stakeholder-specific views, and product-engineering alignment gaps.

## Folder Structure

```text
transcript-intelligence/
  README.md
  requirements.txt
  run_pipeline.py
  build_submission_assets.py
  data/
    raw/dataset/
    processed/
  notebooks/
    01_transcript_intelligence_analysis.ipynb
  src/
    load_data.py
    preprocess.py
    topic_modeling.py
    sentiment.py
    insights.py
    visualization.py
  outputs/
    processed_transcripts.csv
    topic_summary.csv
    sentiment_summary.csv
    additional_insights.csv
    charts/
    slide_assets/
  deck/
    Transcript_Intelligence_Presentation.pptx
  demo_script.md
```

## Setup

```powershell
python -m venv .venv
.\\.venv\\Scripts\\python.exe -m pip install -r requirements.txt
```

## How to Run

```powershell
.\\.venv\\Scripts\\python.exe run_pipeline.py
.\\.venv\\Scripts\\python.exe build_submission_assets.py
```

The default input path is `data/raw/dataset`, copied from the provided assignment folder. You can also pass a different raw folder:

```powershell
.\\.venv\\Scripts\\python.exe run_pipeline.py --raw-dir "path\\to\\dataset"
```

## Methodology

- **Loading:** Recursively discovers transcript folders and reads `meeting-info.json`, `summary.json`, `transcript.json`, `speakers.json`, `speaker-meta.json`, and `events.json`.
- **Normalization:** Produces one row per transcript with meeting metadata, speaker count, transcript text, customer domain, provided topics, action items, key moments, and confidence metadata.
- **Call-type inference:** Uses title/content cues and email domains to infer support, external/customer, or internal calls when no explicit type is present.
- **Topic/theme categorization:** Uses a hybrid approach: TF-IDF clustering for discovery, provided topic metadata for signal, and business keyword rules for human-readable final themes.
- **Sentiment:** Blends the provided summary sentiment score with a transparent local lexical fallback. No paid API key is required.
- **Additional insights:** Detects escalation risk, feature requests, churn/renewal risk, stakeholder dashboard views, and customer-facing vs internal theme alignment.

## Outputs

- `outputs/processed_transcripts.csv`: normalized transcript-level dataset.
- `outputs/topic_summary.csv`: theme names, counts, call type distribution, keywords, examples, and business importance.
- `outputs/sentiment_summary.csv`: sentiment by call type with interpretation.
- `outputs/additional_insights.csv`: combined additional insight output.
- `outputs/charts/*.png`: reusable chart assets for deck and demo.
- `deck/Transcript_Intelligence_Presentation.pptx`: 20-slide leadership presentation.
- `notebooks/01_transcript_intelligence_analysis.ipynb`: readable analysis notebook.

## Limitations

- The taxonomy is intentionally practical, not final. In production, it should be validated with support, product, sales, and engineering stakeholders.
- Sentiment is context-sensitive; outage calls can include constructive language while still representing serious customer risk.
- The source data includes useful summary-level sentiment and topics; the pipeline uses them transparently but can run with local fallbacks.
- Customer/account names are inferred from email domains when explicit account fields are unavailable.

## Future Improvements

- Add optional LLM-assisted theme labeling and excerpt summarization behind an API-key flag.
- Store transcript embeddings for semantic search and similarity-based issue grouping.
- Add CRM/support integrations for account ownership and renewal date context.
- Implement alerting for high-risk escalation/churn transcripts.
- Add human feedback loops to tune taxonomy, sentiment thresholds, and routing ownership.
"""
    README.write_text(readme, encoding="utf-8")

    demo = f"""# Demo Script: Transcript Intelligence

## 1. Introduction

Today I’m walking through Transcript Intelligence, a prototype that turns call transcripts into actionable product, support, account, and engineering insights. The goal is not just NLP output; it is a leadership-ready workflow that helps teams decide what to fix, who should act, and where customer risk is building.

## 2. Folder Structure

I’ll start in the `transcript-intelligence` project folder. The raw JSON transcript folders live in `data/raw/dataset`. The implementation is modular under `src`, with a single `run_pipeline.py` entrypoint. Outputs are written to `outputs`, charts are in `outputs/charts`, and the presentation is in `deck`.

## 3. Run the Pipeline

I’ll run:

```powershell
.\\.venv\\Scripts\\python.exe run_pipeline.py
```

This loads the transcript folders, normalizes one row per transcript, infers call type, computes metadata, assigns themes, scores sentiment, extracts additional insights, and saves the artifacts.

## 4. Processed Transcript Output

Next I’ll open `outputs/processed_transcripts.csv`. The dataset has **{metadata['transcripts']} normalized transcripts** with transcript ID, title, customer domain, call type, speaker count, word count, sentiment score, theme, keywords, and summary fields.

## 5. Topic Categories

Then I’ll show `outputs/topic_summary.csv`. The strongest theme is **{top_theme['theme']}**, with **{int(top_theme['transcript_count'])} transcripts**. The theme summary includes call type distribution, representative keywords, example transcript IDs, and business importance.

## 6. Sentiment Analysis

I’ll show the sentiment charts. The key pattern is that support calls are the most negative, with **{negative_support} negative support transcripts**. That suggests product reliability and support friction are not just operational issues; they are adoption and renewal risks.

## 7. Additional Insights

I’ll show three insight layers:

- Escalation risk detection: flags urgent language, SLA concerns, outage references, and unresolved issues.
- Feature request mining: extracts repeated asks around dashboards, integrations, automation, reports, API, and workflow needs.
- Churn and renewal risk: combines sentiment, renewal language, pricing objections, competitor mentions, and unresolved support issues.

## 8. Recommendations

The practical recommendations are to prioritize the top negative reliability themes, operationalize escalation alerts, route feature requests into roadmap review, and create role-specific dashboards for support, product, account, and engineering leaders.

## 9. Production Vision

In production, this becomes a searchable transcript intelligence layer with batch ingestion first, then real-time risk alerts. It can integrate with CRM, support systems, and product planning tools so that transcript insights move directly into team workflows.
"""
    DEMO.write_text(demo, encoding="utf-8")


def add_textbox(slide, x, y, w, h, text, size=18, color=NAVY, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.TOP
    p = frame.paragraphs[0]
    p.text = str(text)
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = "Aptos"
    if align:
        p.alignment = align
    return box


def add_title(slide, title, subtitle=None):
    add_textbox(slide, 0.6, 0.35, 12.1, 0.5, title, 24, NAVY, True)
    if subtitle:
        add_textbox(slide, 0.62, 0.9, 11.7, 0.3, subtitle, 10.5, GRAY, False)


def add_bullets(slide, x, y, w, h, items, size=15):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = NAVY
        p.font.name = "Aptos"
        p.space_after = Pt(8)
    return box


def add_metric(slide, x, y, value, label, color=BLUE):
    add_textbox(slide, x, y, 2.0, 0.55, value, 34, color, True)
    add_textbox(slide, x, y + 0.55, 2.1, 0.35, label, 9.5, GRAY, False)


def add_table(slide, x, y, w, h, rows, headers):
    table = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h)).table
    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        for p in cell.text_frame.paragraphs:
            p.font.color.rgb = WHITE
            p.font.bold = True
            p.font.size = Pt(9)
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT if r % 2 else WHITE
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(8.2)
                p.font.color.rgb = NAVY
    return table


def add_image(slide, path, x, y, w, h):
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def create_deck():
    processed, topics, sentiment, escalation, features, churn, stakeholder, metadata = load_outputs()
    shutil.copytree(CHARTS, OUTPUTS / "slide_assets", dirs_exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for layout in prs.slide_layouts:
        pass

    # 1
    slide = blank(prs)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    add_textbox(slide, 0.75, 1.15, 11.4, 1.35, "Transcript Intelligence", 48, WHITE, True)
    add_textbox(slide, 0.78, 2.35, 10.2, 0.7, "Turning call transcripts into product, support, and revenue insights", 22, RGBColor(210, 226, 255))
    add_textbox(slide, 0.8, 6.45, 8.5, 0.35, "Take-home assignment | Product and Engineering Leadership", 12, RGBColor(210, 226, 255))
    add_metric(slide, 10.4, 5.25, str(metadata["transcripts"]), "transcripts analyzed", TEAL)
    add_metric(slide, 11.55, 5.25, str(metadata["themes"]), "business themes", AMBER)

    # 2
    slide = blank(prs)
    add_title(slide, "Executive Summary", "The prototype turns raw transcript folders into decision-ready intelligence.")
    add_metric(slide, 0.85, 1.35, str(metadata["transcripts"]), "transcripts")
    add_metric(slide, 3.0, 1.35, str(metadata["call_type_counts"]["support"]), "support calls", CORAL)
    add_metric(slide, 5.15, 1.35, str(metadata["call_type_counts"]["external"]), "external calls", BLUE)
    add_metric(slide, 7.3, 1.35, str(metadata["call_type_counts"]["internal"]), "internal calls", TEAL)
    add_bullets(slide, 0.9, 2.7, 11.4, 2.8, [
        "Support is the clearest risk surface: technical reliability themes carry the lowest average sentiment.",
        "External calls show renewal and expansion context; trust issues cluster around outage, SLA, pricing, and competitive language.",
        "Internal calls expose whether engineering planning is aligned with customer-facing pain.",
        "Recommended path: operationalize escalation risk, route feature requests to PM review, and build stakeholder-specific views."
    ], 15)

    # 3
    slide = blank(prs)
    add_title(slide, "What I Built", "A repeatable batch pipeline with optional room for LLM-assisted labeling later.")
    steps = ["Raw JSON folders", "Normalization", "Call-type inference", "Theme labeling", "Sentiment scoring", "Insight layer", "Stakeholder outputs"]
    x = 0.7
    for i, step in enumerate(steps):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.25), Inches(1.55), Inches(0.85))
        shape.fill.solid(); shape.fill.fore_color.rgb = [BLUE, TEAL, AMBER, CORAL, BLUE, TEAL, NAVY][i]
        shape.line.fill.background()
        tf = shape.text_frame; tf.text = step
        tf.paragraphs[0].font.size = Pt(10); tf.paragraphs[0].font.color.rgb = WHITE; tf.paragraphs[0].font.bold = True
        x += 1.78
    add_bullets(slide, 1.0, 4.0, 10.7, 1.4, [
        "Default path runs locally with pandas, scikit-learn, matplotlib, and transparent rules.",
        "LLM use is deliberately optional, so the project works without paid API keys."
    ], 16)

    # 4
    slide = blank(prs)
    add_title(slide, "Dataset Overview", "The source data includes meeting metadata, transcript turns, topics, summaries, action items, and key moments.")
    add_image(slide, CHARTS / "call_type_breakdown.png", 0.7, 1.25, 5.9, 3.45)
    add_bullets(slide, 7.1, 1.35, 5.2, 3.5, [
        f"{metadata['transcripts']} transcript folders discovered; each contains six JSON files.",
        "No explicit call_type field was present, so type was inferred from title, content, and email domains.",
        "Customer/account identity is inferred from non-aegiscloud.com email domains where available.",
        "The data quality is strong: structured summaries, topics, action items, speaker names, and sentence-level timestamps are available."
    ], 14)

    # 5
    slide = blank(prs)
    add_title(slide, "Topic Categorization Approach", "Hybrid beats pure rules or pure black-box classification for this assignment.")
    add_bullets(slide, 0.9, 1.35, 11.5, 4.5, [
        "Discovery: TF-IDF + clustering reveals natural groupings across titles, summaries, and provided topics.",
        "Interpretation: keyword and phrase analysis converts clusters into human-readable business themes.",
        "Grounding: examples and representative keywords are saved for each theme so leaders can inspect the evidence.",
        "Production path: optional LLM labeling can improve nuance, but should be validated against a stakeholder-approved taxonomy."
    ], 16)

    # 6
    slide = blank(prs)
    add_title(slide, "Themes Identified", "The taxonomy prioritizes business actionability over academic topic labels.")
    rows = []
    for _, r in topics.head(6).iterrows():
        rows.append([r["theme"], int(r["transcript_count"]), r["dominant_call_type"], r["business_importance"][:78]])
    add_table(slide, 0.45, 1.25, 12.45, 5.4, rows, ["Theme", "Count", "Dominant type", "Why it matters"])

    # 7
    slide = blank(prs)
    add_title(slide, "Theme Examples", "Representative summaries show why the labels matter.")
    rows = []
    for _, r in topics.head(3).iterrows():
        ex = str(r["example_transcripts"]).split("|")[0][:145]
        rows.append([r["theme"], ex])
    add_table(slide, 0.55, 1.25, 12.2, 4.8, rows, ["Theme", "Example transcript signal"])

    # 8
    slide = blank(prs)
    add_title(slide, "Sentiment Across Call Types", "Support calls are the strongest negative signal.")
    add_image(slide, CHARTS / "sentiment_distribution_by_call_type.png", 0.7, 1.25, 6.4, 3.7)
    add_bullets(slide, 7.35, 1.45, 5.0, 3.6, [
        "Support has the highest concentration of negative transcripts.",
        "External calls skew positive, but negative moments often carry revenue implications.",
        "Internal calls are mixed: useful for tracking execution pressure and confidence."
    ], 15)

    # 9
    slide = blank(prs)
    add_title(slide, "Sentiment by Theme", "Technical reliability is the most important pain cluster.")
    add_image(slide, CHARTS / "sentiment_by_theme.png", 0.65, 1.15, 7.0, 4.25)
    add_bullets(slide, 8.0, 1.45, 4.7, 3.3, [
        "Reliability and outage themes are negative because customers experience them as lost trust, not just technical inconvenience.",
        "Compliance themes are more positive because the product directly supports urgent audit and regulated-buyer needs."
    ], 15)

    # 10
    slide = blank(prs)
    add_title(slide, "Trend 1: Reliability Is a Product Risk", "Support calls concentrate negative sentiment around outage, SLA, and technical issue language.")
    add_bullets(slide, 0.9, 1.35, 11.2, 4.6, [
        "What it means: technical friction is creating both support burden and account trust risk.",
        "Why it may be happening: customers depend on Detect/backup/security workflows where downtime has high business impact.",
        "Who should care: support leaders, product managers, and engineering leads.",
        "Action: prioritize reliability fixes, publish clearer incident comms, and create an escalation review for top negative reliability transcripts."
    ], 16)

    # 11
    slide = blank(prs)
    add_title(slide, "Trend 2: External Calls Carry Renewal Context", "External conversations reveal both expansion opportunity and risk signals.")
    add_bullets(slide, 0.9, 1.35, 11.1, 4.6, [
        "What it means: positive external sentiment is encouraging, but risk terms around renewal, pricing, SLA, and competitors need account workflows.",
        "Why it may be happening: buyers see value in compliance and security outcomes, while reliability incidents affect confidence.",
        "Who should care: sales leaders, account managers, and customer success.",
        "Action: connect risk scores to CRM ownership and trigger proactive renewal plays."
    ], 16)

    # 12
    slide = blank(prs)
    add_title(slide, "Additional Insight: Escalation Risk", "Urgent language and SLA/outage markers create an early-warning queue.")
    rows = []
    for _, r in escalation.head(5).iterrows():
        rows.append([r["transcript_id"], r["call_type"], int(r["risk_score"]), r["suggested_owner"], r["reason"][:50]])
    add_table(slide, 0.55, 1.25, 12.2, 4.4, rows, ["Transcript", "Type", "Risk", "Owner", "Reason"])
    add_textbox(slide, 0.75, 6.05, 11.3, 0.45, "Value: helps teams catch customer pain before it becomes churn, SLA exposure, or executive escalation.", 14, NAVY, True)

    # 13
    slide = blank(prs)
    add_title(slide, "Additional Insight: Feature Request Mining", "Customer language becomes roadmap evidence.")
    rows = []
    for _, r in features.head(6).iterrows():
        rows.append([r["requested_capability"], int(r["frequency"]), str(r["impacted_call_types"])[:52]])
    add_table(slide, 0.65, 1.3, 7.2, 4.4, rows, ["Requested capability", "Frequency", "Call types"])
    add_bullets(slide, 8.25, 1.45, 4.3, 3.8, [
        "Product managers can inspect examples rather than rely on anecdotes.",
        "Repeated asks around integrations, API, dashboards, automation, and reporting become roadmap candidates.",
        "Pair frequency with revenue/risk context before prioritizing."
    ], 14)

    # 14
    slide = blank(prs)
    add_title(slide, "Additional Insight: Churn and Stakeholder Views", "Different teams need different slices of the same transcript intelligence.")
    rows = []
    for _, r in stakeholder.iterrows():
        rows.append([r["stakeholder"], r["top_relevant_theme"], int(r["negative_transcripts"]), r["recommended_action"]])
    add_table(slide, 0.5, 1.25, 12.3, 4.9, rows, ["Stakeholder", "Top theme", "Neg.", "Recommended action"])

    # 15
    slide = blank(prs)
    add_title(slide, "Product Vision", "From analysis prototype to enterprise transcript intelligence layer.")
    add_bullets(slide, 0.9, 1.35, 11.1, 4.8, [
        "Searchable transcript intelligence across support, external, and internal calls.",
        "Topic and sentiment dashboards for support, product, account, and engineering leaders.",
        "Risk alerts for escalations, renewal threats, and repeated unresolved issues.",
        "Voice-of-customer mining that feeds product roadmap and documentation priorities.",
        "Integrations with CRM, support tools, and project planning systems."
    ], 16)

    # 16
    slide = blank(prs)
    add_title(slide, "Technical Architecture", "Start batch-first; evolve toward near-real-time routing and alerting.")
    steps = ["Ingestion", "Preprocess", "Embeddings / TF-IDF", "Classify + sentiment", "Insight generation", "Dashboard / API"]
    x = 0.8
    for step in steps:
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(2.45), Inches(1.75), Inches(0.85))
        shape.fill.solid(); shape.fill.fore_color.rgb = LIGHT
        shape.line.color.rgb = BLUE
        shape.text_frame.text = step
        shape.text_frame.paragraphs[0].font.size = Pt(11)
        shape.text_frame.paragraphs[0].font.color.rgb = NAVY
        shape.text_frame.paragraphs[0].font.bold = True
        x += 2.05
    add_bullets(slide, 1.0, 4.15, 10.8, 1.4, [
        "The current implementation is deterministic and inspectable; LLM steps can be added for labeling and summarization when validation is available.",
        "A production system should store transcript metadata, embeddings, predictions, examples, and human feedback separately."
    ], 14)

    # 17
    slide = blank(prs)
    add_title(slide, "Limitations and Tradeoffs", "The design is intentionally practical and transparent.")
    add_bullets(slide, 0.9, 1.25, 11.4, 5.0, [
        "Sample size is modest; themes should be validated as more transcripts arrive.",
        "Call type is inferred because no explicit field exists.",
        "Sentiment is context-sensitive; constructive outage calls can still indicate serious customer risk.",
        "Rule-based taxonomy is explainable but may miss nuance; LLM labeling would need evaluation and review.",
        "Account names are inferred from email domains when explicit customer fields are unavailable."
    ], 16)

    # 18
    slide = blank(prs)
    add_title(slide, "Recommendations", "Make the prototype useful by connecting insights to ownership.")
    add_bullets(slide, 0.9, 1.3, 11.2, 4.9, [
        "Prioritize the top negative reliability themes with support and engineering.",
        "Create a weekly escalation-risk workflow with owners and SLA follow-through.",
        "Use feature request mining as an input to roadmap review, not as an automatic prioritization score.",
        "Build stakeholder dashboards around the questions each team actually asks.",
        "Validate taxonomy and thresholds with support, product, sales, and engineering leaders."
    ], 16)

    # 19
    slide = blank(prs)
    add_title(slide, "Demo Walkthrough", "The demo is designed to be short, concrete, and reproducible.")
    add_bullets(slide, 0.9, 1.3, 11.3, 4.8, [
        "Show the project structure and raw dataset folder.",
        "Run `run_pipeline.py` and confirm outputs are generated.",
        "Open processed transcripts, topic summary, sentiment summary, and additional insights.",
        "Show charts used by the deck.",
        "Close with productization path: search, dashboards, risk alerts, and workflow integrations."
    ], 16)

    # 20
    slide = blank(prs)
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = NAVY
    add_textbox(slide, 0.8, 2.55, 11.6, 0.95, "Q&A", 54, WHITE, True, PP_ALIGN.CENTER)
    add_textbox(slide, 2.0, 3.45, 9.4, 0.55, "How would you use transcripts to change decisions before problems become obvious?", 18, RGBColor(210, 226, 255), False, PP_ALIGN.CENTER)

    DECK.parent.mkdir(parents=True, exist_ok=True)
    prs.save(DECK)


def create_notebook():
    processed, topics, sentiment, escalation, features, churn, stakeholder, metadata = load_outputs()
    nb = nbf.v4.new_notebook()
    nb["cells"] = [
        nbf.v4.new_markdown_cell("# Transcript Intelligence Analysis\n\nThis notebook documents the end-to-end analysis for the take-home assignment. It is written to be readable by product and engineering leadership as well as technical reviewers."),
        nbf.v4.new_markdown_cell(f"## 1. Assignment Objective\n\nBuild a pipeline that processes approximately 100 call transcripts across support, external/customer, and internal calls; categorize them into themes; analyze sentiment; and identify additional product/business insights. This run processed **{metadata['transcripts']} transcripts**."),
        nbf.v4.new_code_cell("from pathlib import Path\nimport pandas as pd\nROOT = Path.cwd()\nif not (ROOT / 'outputs').exists() and ROOT.name == 'notebooks':\n    ROOT = ROOT.parent\nprocessed = pd.read_csv(ROOT / 'outputs' / 'processed_transcripts.csv')\ntopics = pd.read_csv(ROOT / 'outputs' / 'topic_summary.csv')\nsentiment = pd.read_csv(ROOT / 'outputs' / 'sentiment_summary.csv')\nescalation = pd.read_csv(ROOT / 'outputs' / 'escalation_risks.csv')\nfeatures = pd.read_csv(ROOT / 'outputs' / 'feature_requests.csv')\nchurn = pd.read_csv(ROOT / 'outputs' / 'churn_risks.csv')\nprocessed.head()"),
        nbf.v4.new_markdown_cell("## 2. Dataset Overview\n\nThe raw data is a nested JSON folder structure. Each transcript folder contains meeting metadata, sentence-level transcript rows, speaker information, event rows, summary, action items, topics, sentiment, and key moments. Because no explicit `call_type` field exists, the pipeline infers support/external/internal using meeting title, content cues, and email domains."),
        nbf.v4.new_code_cell("processed[['transcript_id','title','call_type','customer_domain','speaker_count','word_count','theme','sentiment_label']].head(10)"),
        nbf.v4.new_code_cell("processed['call_type'].value_counts()"),
        nbf.v4.new_markdown_cell("![Call type breakdown](../outputs/charts/call_type_breakdown.png)\n\n**Call-type inference result:** the corpus resolves to 50 support calls, 35 external/customer calls, and 15 internal calls. This matters because each stakeholder group reads transcript intelligence through a different lens: support needs issue triage, account teams need renewal risk, and engineering needs recurring technical patterns."),
        nbf.v4.new_markdown_cell("## 3. Data Cleaning and Assumptions\n\nCleaning removes obvious whitespace/timestamp noise while preserving the transcript language. Customer/account names are inferred from non-`aegiscloud.com` email domains. This is clearly documented as an assumption rather than treated as a perfect CRM account mapping."),
        nbf.v4.new_markdown_cell("## 4. Topic / Theme Categorization Approach\n\nThe pipeline uses a hybrid method: TF-IDF clustering for discovery, provided topic metadata for grounding, and business keyword rules for final human-readable themes. This is more practical than pure clustering, which often produces hard-to-explain labels, and safer than pure LLM labeling, which would require API access and validation."),
        nbf.v4.new_code_cell("topics[['theme','transcript_count','dominant_call_type','avg_sentiment_score','representative_keywords','business_importance']]"),
        nbf.v4.new_markdown_cell("![Themes identified](../outputs/charts/theme_counts.png)\n\nThe taxonomy is intentionally business-readable. For example, **Product bugs and technical reliability** is more useful to leadership than a raw cluster label such as `cluster_2`, because it points directly to support burden, engineering ownership, and adoption risk."),
        nbf.v4.new_markdown_cell("### Example Transcript IDs and Excerpts by Theme\n\nEach theme includes example transcript IDs and short excerpts so a reviewer can inspect whether the category is grounded in the underlying calls."),
        nbf.v4.new_code_cell("topic_examples = topics[['theme','example_transcripts','business_importance']].copy()\ntopic_examples"),
        nbf.v4.new_markdown_cell("## 5. Sentiment Analysis Approach\n\nThe source summaries include sentiment scores. The default implementation blends those structured scores with a transparent local lexical fallback. This keeps the pipeline runnable without paid API keys while making the scoring logic inspectable."),
        nbf.v4.new_code_cell("sentiment"),
        nbf.v4.new_markdown_cell("![Sentiment distribution by call type](../outputs/charts/sentiment_distribution_by_call_type.png)\n\n![Average sentiment by call type](../outputs/charts/average_sentiment_by_call_type.png)\n\n**Interpretation:** support calls have the clearest negative signal. That is expected because customers usually contact support when something is broken, but the concentration around reliability/outage/SLA language means the signal should not stay inside support operations. Product and engineering should treat it as adoption and renewal risk."),
        nbf.v4.new_markdown_cell("## 6. Insight Findings\n\nThe most important trend is that support calls show the strongest negative sentiment, especially in reliability/outage/SLA-related themes. This should matter to support, product, and engineering because technical friction is showing up as customer trust risk, not just ticket volume.\n\n![Sentiment by theme](../outputs/charts/sentiment_by_theme.png)\n\n![Theme x call type sentiment heatmap](../outputs/charts/theme_calltype_sentiment_heatmap.png)"),
        nbf.v4.new_code_cell("topics.sort_values('avg_sentiment_score')[['theme','transcript_count','dominant_call_type','avg_sentiment_score','business_importance']]"),
        nbf.v4.new_markdown_cell("## 7. Additional Insights\n\nThe assignment asks, 'What else can you see?' The project implements escalation risk detection, feature request mining, churn/renewal risk signals, stakeholder-specific views, and product-engineering feedback-loop alignment. These are designed as product surfaces, not just analysis tables: each insight has a likely owner and suggested action."),
        nbf.v4.new_code_cell("escalation[['transcript_id','call_type','theme','risk_score','reason','suggested_owner']].head(10)"),
        nbf.v4.new_code_cell("features.head(10)"),
        nbf.v4.new_code_cell("churn[['transcript_id','customer_or_account','risk_score','risk_terms','suggested_action']].head(10)"),
        nbf.v4.new_code_cell("stakeholder = pd.read_csv(ROOT / 'outputs' / 'stakeholder_views.csv')\nstakeholder"),
        nbf.v4.new_code_cell("pe_gap = pd.read_csv(ROOT / 'outputs' / 'product_engineering_gap.csv')\npe_gap"),
        nbf.v4.new_markdown_cell("## 8. Limitations\n\nThe taxonomy should be validated with stakeholders. Sentiment is context-sensitive, especially for incident calls. Account names are inferred from domains, not joined to CRM. Optional LLM labeling would add nuance, but it should be treated as an assisted workflow with human review."),
        nbf.v4.new_markdown_cell("## 9. Recommendations\n\nPrioritize negative reliability themes, create an escalation-risk workflow, route feature requests into roadmap review, build stakeholder-specific dashboards, and validate the taxonomy with product/support/sales/engineering leadership."),
        nbf.v4.new_markdown_cell("## 10. Productionization Next Steps\n\nMove from batch CSV outputs to a service-backed transcript intelligence layer: ingestion jobs, metadata store, embeddings index, validated classifiers, alerting rules, CRM/support integrations, and human feedback loops for taxonomy and model quality."),
    ]
    NOTEBOOK.parent.mkdir(parents=True, exist_ok=True)
    nbf.write(nb, NOTEBOOK)


def main():
    write_markdown_assets()
    create_notebook()
    create_deck()
    print(f"README: {README}")
    print(f"Demo script: {DEMO}")
    print(f"Notebook: {NOTEBOOK}")
    print(f"Deck: {DECK}")


if __name__ == "__main__":
    main()
